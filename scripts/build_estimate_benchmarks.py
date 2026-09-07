from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook

from app.estimate_loader import EstimateGroup, EstimateLoader, EstimateRoom


OUTPUT_PATH = Path("data/cleaned/estimate_benchmarks.json")
NON_ROOM_TYPES = {"general_work", "plumbing", "electrical", "heating"}


def most_repeated_area(values: list[float]) -> float:
    counts = Counter(values)
    max_count = max(counts.values())
    return max(value for value, count in counts.items() if count == max_count)


def infer_object_area_m2(rooms: list[EstimateRoom]) -> float | None:
    shared_candidates = [
        line.quantity
        for room in rooms
        if room.room_type in NON_ROOM_TYPES
        for group in room.groups
        if group.work_area_code == "floor"
        for line in group.lines
        if line.unit == "м²" and line.quantity and line.quantity > 0
    ]
    if shared_candidates:
        return round(most_repeated_area(shared_candidates), 2)

    room_areas: list[float] = []
    for room in rooms:
        if room.room_type in NON_ROOM_TYPES:
            continue

        candidates = [
            line.quantity
            for group in room.groups
            if group.work_area_code == "floor"
            for line in group.lines
            if line.unit == "м²" and line.quantity and line.quantity > 0
        ]
        if candidates:
            room_areas.append(most_repeated_area(candidates))

    if not room_areas:
        return None
    return round(sum(room_areas), 2)


def stage_scope(
    grouped: list[tuple[EstimateRoom, EstimateGroup]],
    physical_rooms: list[EstimateRoom],
) -> str:
    service_names = [
        line.service_name.lower().replace("ё", "е")
        for _, group in grouped
        for line in group.lines
    ]
    if any("по квартире" in name or "всей квартире" in name for name in service_names):
        return "full_object"

    stage_room_ids = {
        id(room)
        for room, _ in grouped
        if room.room_type not in NON_ROOM_TYPES
    }
    work_areas = {group.work_area_code for _, group in grouped}
    required_rooms = max(2, round(len(physical_rooms) * 0.5))

    if len(stage_room_ids) >= required_rooms and len(work_areas) >= 3:
        return "full_object"
    return "partial"


def build_benchmarks(estimates_dir: Path) -> dict:
    loader = EstimateLoader(estimates_dir)
    benchmarks: list[dict] = []

    for path in sorted(estimates_dir.glob("*.xlsx")):
        workbook = load_workbook(path, data_only=True, read_only=True)
        if loader.sheet_name not in workbook.sheetnames:
            workbook.close()
            continue

        rows = list(
            workbook[loader.sheet_name].iter_rows(
                min_col=1,
                max_col=7,
                values_only=True,
            )
        )
        workbook.close()

        rooms = loader._parse_rooms(rows)
        object_area_m2 = infer_object_area_m2(rooms)
        if object_area_m2 is None:
            continue

        grouped_by_stage: dict[str, list[tuple[EstimateRoom, EstimateGroup]]] = defaultdict(list)
        for room in rooms:
            for group in room.groups:
                grouped_by_stage[group.stage_code].append((room, group))

        physical_rooms = [room for room in rooms if room.room_type not in NON_ROOM_TYPES]
        stage_codes = set(grouped_by_stage)
        required_full_stages = {"preparation", "rough_finish", "finish"}
        object_scope = (
            "full_object"
            if len(physical_rooms) >= 2 and required_full_stages.issubset(stage_codes)
            else "partial"
        )
        object_total = round(
            sum(
                room.total_after_discount
                if room.total_after_discount is not None
                else room.total_before_discount
                if room.total_before_discount is not None
                else sum(line.line_total for group in room.groups for line in group.lines)
                for room in rooms
            ),
            2,
        )
        stages: list[dict] = []

        for stage_code, grouped in sorted(grouped_by_stage.items()):
            lines = [line for _, group in grouped for line in group.lines]
            total = round(sum(line.line_total for line in lines), 2)
            stage_name = Counter(group.stage_name for _, group in grouped).most_common(1)[0][0]
            work_area_codes = sorted({
                group.work_area_code
                for _, group in grouped
                if group.work_area_code and group.work_area_code != "general"
            })
            work_items = sorted({
                line.service_name
                for line in lines
                if line.service_name
            })
            stages.append(
                {
                    "stage_code": stage_code,
                    "stage_name": stage_name,
                    "scope": stage_scope(grouped, physical_rooms),
                    "total": total,
                    "cost_per_m2": round(total / object_area_m2, 2),
                    "line_items_count": len(lines),
                    "work_area_codes": work_area_codes,
                    "work_items": work_items,
                }
            )

        benchmarks.append(
            {
                "estimate_id": path.stem,
                "object_area_m2": object_area_m2,
                "object_scope": object_scope,
                "object_total": object_total,
                "object_cost_per_m2": round(object_total / object_area_m2, 2),
                "stages": stages,
            }
        )

    return {"version": 1, "benchmarks": benchmarks}


def main() -> None:
    load_dotenv()
    estimates_dir_value = os.getenv("ESTIMATES_DIR")
    if not estimates_dir_value:
        raise ValueError("ESTIMATES_DIR is not set")

    result = build_benchmarks(Path(estimates_dir_value).expanduser())
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(
        json.dumps(result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"objects: {len(result['benchmarks'])}")
    print(f"saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
