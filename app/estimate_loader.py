from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.parser import ParsedDocument


WORK_AREA_LABELS = {
    "пол": "floor",
    "полы": "floor",
    "стена": "walls",
    "стены": "walls",
    "потолок": "ceiling",
    "потолки": "ceiling",
    "откос": "slopes",
    "откосы": "slopes",
    "двери/окна": "doors_windows",
    "двери, окна": "doors_windows",
    "сантехника": "plumbing",
    "электрика": "electrical",
    "прочее": "other",
}

STAGE_LABELS = {
    "демонтажные работы": "demolition",
    "подготовительные работы": "preparation",
    "черновые отделочные работы": "rough_finish",
    "чистовые отделочные работы": "finish",
    "сантехнические работы": "plumbing",
    "электромонтажные работы": "electrical",
    "отопление": "heating",
}

UNIT_LABELS = {
    "м2": "м²",
    "м²": "м²",
    "м.п": "м",
    "м.п.": "м",
    "п.м": "м",
    "п.м.": "м",
    "шт": "шт.",
    "шт.": "шт.",
}

MEASUREMENT_RE = re.compile(r"^\s*(-?\d+(?:[.,]\d+)?)\s*(.*?)\s*$")
DATE_RE = re.compile(r"\b(\d{2})\.(\d{2})\.(\d{4})\b")


@dataclass(slots=True)
class EstimateLine:
    position: int
    service_name: str
    quantity: float | None
    unit: str
    measurement_original: str
    unit_price: float
    line_total: float


@dataclass(slots=True)
class EstimateGroup:
    stage_name: str
    stage_code: str
    work_area_name: str
    work_area_code: str
    lines: list[EstimateLine] = field(default_factory=list)


@dataclass(slots=True)
class EstimateRoom:
    name: str
    room_type: str
    groups: list[EstimateGroup]
    total_before_discount: float | None
    total_after_discount: float | None


class EstimateLoader:
    """Convert company estimate workbooks into anonymized RAG documents."""

    def __init__(self, estimates_dir: str | Path, sheet_name: str = "Смета") -> None:
        self.estimates_dir = Path(estimates_dir).expanduser()
        self.sheet_name = sheet_name

    def load_documents(self) -> list[ParsedDocument]:
        if not self.estimates_dir.is_dir():
            raise FileNotFoundError(f"Estimates directory not found: {self.estimates_dir}")

        documents: list[ParsedDocument] = []
        for path in sorted(self.estimates_dir.glob("*.xlsx")):
            documents.extend(self.load_workbook_documents(path))

        return documents

    def load_workbook_documents(self, path: str | Path) -> list[ParsedDocument]:
        workbook_path = Path(path)
        workbook = load_workbook(
            workbook_path,
            data_only=True,
            read_only=True,
        )

        if self.sheet_name not in workbook.sheetnames:
            workbook.close()
            return []

        worksheet = workbook[self.sheet_name]
        rows = list(
            worksheet.iter_rows(
                min_col=1,
                max_col=7,
                values_only=True,
            )
        )
        workbook.close()

        estimate_id = workbook_path.stem
        estimate_date = self._extract_estimate_date(rows)
        rooms = self._parse_rooms(rows)

        documents: list[ParsedDocument] = []
        for room_index, room in enumerate(rooms, start=1):
            documents.append(
                self._build_room_summary_document(
                    estimate_id=estimate_id,
                    source_file=workbook_path.name,
                    estimate_date=estimate_date,
                    room_index=room_index,
                    room=room,
                )
            )

            for group_index, group in enumerate(room.groups, start=1):
                documents.append(
                    self._build_group_document(
                        estimate_id=estimate_id,
                        source_file=workbook_path.name,
                        estimate_date=estimate_date,
                        room_index=room_index,
                        group_index=group_index,
                        room=room,
                        group=group,
                    )
                )

        return documents

    def _parse_rooms(self, rows: list[tuple]) -> list[EstimateRoom]:
        room_starts = self._find_room_starts(rows)
        rooms: list[EstimateRoom] = []

        for room_number, (start_index, room_name) in enumerate(room_starts):
            end_index = (
                room_starts[room_number + 1][0]
                if room_number + 1 < len(room_starts)
                else len(rows)
            )
            room = self._parse_room_block(
                room_name=room_name,
                rows=rows[start_index:end_index],
            )
            if room.groups:
                rooms.append(room)

        return rooms

    def _find_room_starts(self, rows: list[tuple]) -> list[tuple[int, str]]:
        starts: list[tuple[int, str]] = []

        for index, row in enumerate(rows):
            first = self._cell(row, 0)
            second = self._clean_text(self._cell(row, 1))
            if first not in (None, "") or not second:
                continue

            lookahead = rows[index + 1:index + 6]
            if any(self._is_table_header(candidate) for candidate in lookahead):
                starts.append((index, second))

        return starts

    def _parse_room_block(self, room_name: str, rows: list[tuple]) -> EstimateRoom:
        groups: dict[tuple[str, str, str, str], EstimateGroup] = {}
        current_stage_name = "Без этапа"
        current_stage_code = "unspecified"
        current_work_area_name = "Общие работы"
        current_work_area_code = "general"
        total_before_discount: float | None = None
        total_after_discount: float | None = None

        for row in rows:
            first = self._cell(row, 0)
            service_name = self._clean_text(self._cell(row, 1))
            measurement = self._cell(row, 4)
            unit_price = self._as_float(self._cell(row, 5))
            line_total = self._as_float(self._cell(row, 6))

            if isinstance(first, str):
                label = self._clean_text(first)
                normalized_label = self._normalize_label(label)

                if normalized_label.startswith("итого по стоимости работ со скидкой"):
                    total_after_discount = line_total
                    continue

                if normalized_label.startswith("итого по стоимости работ"):
                    total_before_discount = line_total
                    continue

                if line_total is None:
                    continue

                work_area_code = WORK_AREA_LABELS.get(normalized_label)
                if work_area_code:
                    current_work_area_name = label
                    current_work_area_code = work_area_code
                    continue

                current_stage_name = label
                current_stage_code = STAGE_LABELS.get(normalized_label, "other")
                current_work_area_name = "Общие работы"
                current_work_area_code = "general"
                continue

            if not isinstance(first, (int, float)) or not service_name:
                continue
            if unit_price is None or line_total is None:
                continue

            quantity, unit, measurement_original = self._parse_measurement(measurement)
            key = (
                current_stage_name,
                current_stage_code,
                current_work_area_name,
                current_work_area_code,
            )
            group = groups.setdefault(
                key,
                EstimateGroup(
                    stage_name=current_stage_name,
                    stage_code=current_stage_code,
                    work_area_name=current_work_area_name,
                    work_area_code=current_work_area_code,
                ),
            )
            group.lines.append(
                EstimateLine(
                    position=int(first),
                    service_name=service_name,
                    quantity=quantity,
                    unit=unit,
                    measurement_original=measurement_original,
                    unit_price=unit_price,
                    line_total=line_total,
                )
            )

        return EstimateRoom(
            name=room_name,
            room_type=self._normalize_room_type(room_name),
            groups=list(groups.values()),
            total_before_discount=total_before_discount,
            total_after_discount=total_after_discount,
        )

    def _build_room_summary_document(
        self,
        estimate_id: str,
        source_file: str,
        estimate_date: str | None,
        room_index: int,
        room: EstimateRoom,
    ) -> ParsedDocument:
        calculated_total = round(
            sum(line.line_total for group in room.groups for line in group.lines),
            2,
        )
        stage_totals: dict[str, float] = defaultdict(float)
        for group in room.groups:
            stage_totals[group.stage_name] += sum(line.line_total for line in group.lines)

        lines = [
            "Тип данных: исторический пример реальной сметы.",
            f"Помещение или раздел: {room.name}",
            f"Нормализованный тип: {room.room_type}",
        ]
        if estimate_date:
            lines.append(f"Дата сметы: {estimate_date}")
        if room.total_before_discount is not None:
            lines.append(f"Итого работ до скидки: {room.total_before_discount:.2f} ₽")
        if room.total_after_discount is not None:
            lines.append(f"Итого работ после скидки: {room.total_after_discount:.2f} ₽")
        if room.total_before_discount is None:
            lines.append(f"Сумма позиций: {calculated_total:.2f} ₽")

        lines.append("Этапы работ:")
        for stage_name, stage_total in stage_totals.items():
            lines.append(f"- {stage_name}: {stage_total:.2f} ₽")

        lines.append(
            "Примечание: это смета конкретного объекта, а не действующий прайс или гарантированная цена."
        )

        title = f"Реальная смета: {room.name}"
        url = f"estimate://{estimate_id}/room/{room_index}/summary"
        metadata = self._base_metadata(
            estimate_id=estimate_id,
            source_file=source_file,
            estimate_date=estimate_date,
            room_index=room_index,
            room=room,
        )
        metadata.update(
            {
                "estimate_document_type": "room_summary",
                "work_total": room.total_after_discount
                or room.total_before_discount
                or calculated_total,
                "work_total_before_discount": room.total_before_discount,
                "work_total_after_discount": room.total_after_discount,
                "line_items_count": sum(len(group.lines) for group in room.groups),
            }
        )

        return ParsedDocument(
            url=url,
            title=title,
            h1=title,
            text="\n".join(lines),
            headings=[room.name, "Этапы работ"],
            metadata=metadata,
        )

    def _build_group_document(
        self,
        estimate_id: str,
        source_file: str,
        estimate_date: str | None,
        room_index: int,
        group_index: int,
        room: EstimateRoom,
        group: EstimateGroup,
    ) -> ParsedDocument:
        group_total = round(sum(line.line_total for line in group.lines), 2)
        lines = [
            "Тип данных: исторический пример реальной сметы.",
            f"Помещение или раздел: {room.name}",
            f"Этап: {group.stage_name}",
            f"Раздел работ: {group.work_area_name}",
            "Работы:",
        ]

        for item in group.lines:
            measurement = item.measurement_original or "количество не указано"
            lines.append(
                f"- {item.service_name}: {measurement}; "
                f"{item.unit_price:.2f} ₽ за единицу; итого {item.line_total:.2f} ₽"
            )

        lines.extend(
            [
                f"Итого по группе: {group_total:.2f} ₽",
                "Примечание: это смета конкретного объекта, а не действующий прайс или гарантированная цена.",
            ]
        )

        title = f"{room.name}: {group.stage_name}, {group.work_area_name}"
        url = f"estimate://{estimate_id}/room/{room_index}/group/{group_index}"
        metadata = self._base_metadata(
            estimate_id=estimate_id,
            source_file=source_file,
            estimate_date=estimate_date,
            room_index=room_index,
            room=room,
        )
        metadata.update(
            {
                "estimate_document_type": "work_group",
                "stage_name": group.stage_name,
                "stage_code": group.stage_code,
                "work_area_name": group.work_area_name,
                "work_area_code": group.work_area_code,
                "group_total": group_total,
                "line_items_count": len(group.lines),
            }
        )

        return ParsedDocument(
            url=url,
            title=title,
            h1=title,
            text="\n".join(lines),
            headings=[room.name, group.stage_name, group.work_area_name],
            metadata=metadata,
        )

    def _base_metadata(
        self,
        estimate_id: str,
        source_file: str,
        estimate_date: str | None,
        room_index: int,
        room: EstimateRoom,
    ) -> dict:
        return {
            "document_type": "estimate",
            "source_type": "real_estimate",
            "source_file": source_file,
            "estimate_id": estimate_id,
            "estimate_date": estimate_date,
            "room_index": room_index,
            "room_name": room.name,
            "room_type": room.room_type,
            "currency": "RUB",
            "is_historical_example": True,
        }

    def _parse_measurement(self, value) -> tuple[float | None, str, str]:
        original = self._clean_text(value)
        if not original:
            return None, "", ""

        match = MEASUREMENT_RE.match(original)
        if not match:
            return None, "", original

        quantity = float(match.group(1).replace(",", "."))
        raw_unit = self._normalize_label(match.group(2))
        unit = UNIT_LABELS.get(raw_unit, match.group(2).strip())
        return quantity, unit, original

    def _extract_estimate_date(self, rows: list[tuple]) -> str | None:
        for row in rows[:10]:
            for value in row:
                if isinstance(value, (datetime, date)):
                    return value.isoformat()
                if not isinstance(value, str):
                    continue
                match = DATE_RE.search(value)
                if match:
                    day, month, year = match.groups()
                    return f"{year}-{month}-{day}"
        return None

    def _normalize_room_type(self, room_name: str) -> str:
        normalized = self._normalize_label(room_name)

        if "общ" in normalized and "работ" in normalized:
            return "general_work"
        if "элект" in normalized:
            return "electrical"
        if "сантех" in normalized:
            return "plumbing"
        if "отоп" in normalized:
            return "heating"
        if "детск" in normalized:
            return "children_room"
        if "спаль" in normalized:
            return "bedroom"
        if "ванн" in normalized:
            return "bathroom"
        if (
                "сануз" in normalized
                or "с/у" in normalized
                or "душев" in normalized
                or re.search(r"(^|[\s-])су($|[\s-])", normalized)
        ):
            return "bathroom"
        if "кух" in normalized and "гост" in normalized:
            return "kitchen_living_room"
        if "кух" in normalized:
            return "kitchen"
        if "гост" in normalized:
            return "living_room"
        if any(token in normalized for token in ("коридор", "прихож")):
            return "hallway"
        if "гардероб" in normalized:
            return "wardrobe"
        if "лоджи" in normalized or "балкон" in normalized:
            return "balcony"
        if "кабинет" in normalized:
            return "office"
        if "постир" in normalized:
            return "laundry"
        if "комнат" in normalized:
            return "room"
        return "other"

    def _is_table_header(self, row: tuple) -> bool:
        first = self._normalize_label(self._cell(row, 0))
        second = self._normalize_label(self._cell(row, 1))
        return first.startswith("№ п.п") and second == "наименование"

    @staticmethod
    def _cell(row: tuple, index: int):
        return row[index] if index < len(row) else None

    @staticmethod
    def _as_float(value) -> float | None:
        if isinstance(value, (int, float)):
            return float(value)
        return None

    @staticmethod
    def _clean_text(value) -> str:
        if value is None:
            return ""
        return " ".join(str(value).replace("\xa0", " ").split()).strip()

    @classmethod
    def _normalize_label(cls, value) -> str:
        text = cls._clean_text(value).lower().replace("ё", "е")
        return re.sub(r"\s+", " ", text).strip(" :")
